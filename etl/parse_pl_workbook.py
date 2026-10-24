"""
Parse the firm's P&L (income statement) Excel workbook into a tidy table
of actuals by line item and month. This is the baseline for the Budget
page's overhead lines (prior year + % change).

Not built from a real sample yet, so it reads the common layouts rather
than one fixed shape:

- **Columns.** A header row with month columns ("Jan 2026", "January
  2026", "Jan-26", "1/2026", "2026-01", real dates, or bare "Jan" with
  the year taken from the report title). If there are no month columns, a
  single "Total" / "YTD" / "Annual" / "2025" column is read as a full-year
  amount. A trailing Total column next to month columns is ignored.
- **Rows.** One line item per row: "6100 · Rent", "6100 Rent" or just
  "Rent". Labels can sit in one column with leading spaces, or step
  across columns (QuickBooks indents by column). Rows with a label and no
  amounts are headers: "Income" / "Revenue", "Cost of Goods Sold",
  "Expenses", "Other Income", "Other Expense" set the section; anything
  else becomes the category for the rows under it. Subtotals ("Total
  ...", "Gross Profit", "Net Income", ...) are skipped so nothing is
  counted twice.
- **Sheets.** Every sheet that has line items is read, one entity per
  sheet (from the title rows, or the sheet name). A sheet named like
  "Consolidated" / "Combined" / "Summary" is flagged so the Budget page
  doesn't add it on top of the entity sheets.

Amounts are kept as shown: revenue and expenses are both positive on a
normal P&L. If no section header is found, the account number's leading
digit decides (config.GL_ACCOUNT_TYPE_PREFIXES).

Output grain: one row per (sheet, line item, period):
    entity, sheet, is_consolidated, section, category, account_number,
    line_item, period_start, period_end, months, amount, source_file
"""
from __future__ import annotations

import calendar
import csv
import logging
import re
import sys
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl

import config
from etl.common import find_all_files, parse_money

logger = logging.getLogger("somos.etl.pl_workbook")

_MONTHS = {m.lower(): i for i, m in enumerate(calendar.month_abbr) if m}
_MONTHS.update({m.lower(): i for i, m in enumerate(calendar.month_name) if m})
_MONTHS["sept"] = 9

# "Jan 2026", "January 2026", "Jan-26", "Jan '26", "Jan 26", "Jan"
_MONTH_NAME_RE = re.compile(r"^([A-Za-z]{3,9})\.?(?:[\s\-/']+'?(\d{2}|\d{4}))?$")
# "1/2026", "01-2026", "2026-01", "2026/1"
_MONTH_NUM_RE = re.compile(r"^(?:(\d{1,2})[/\-](\d{4})|(\d{4})[/\-](\d{1,2}))$")
_YEAR_RE = re.compile(r"\b(20\d{2})\b")
_TOTAL_HEADER_RE = re.compile(r"^(total|ytd|year[\s\-]*to[\s\-]*date|annual|full\s*year|fy\s*\d{2,4}|\d{4})\b", re.I)
_ACCOUNT_RE = re.compile(r"^(\d{3,6}(?:[.\-]\d{1,4})?)\s*(?:[·•:\-–]\s*|\s+)(.+)$")
_ENTITY_HINT_RE = re.compile(r"\b(llc|llp|inc|corp|group|pllc|pc|ltd|company)\b", re.I)
_CONSOLIDATED_RE = re.compile(r"consolid|combined|summary|all\s+entities|total", re.I)

# Subtotal / computed rows -- skipped so line items aren't counted twice.
_SKIP_RE = re.compile(
    r"^(total\b|gross\s+(profit|margin)|net\s+(income|profit|loss|ordinary|operating|other)|"
    r"operating\s+(income|profit)|ebitda|ebit\b|income\s+before|profit\s+before)",
    re.I,
)

# Section headers, checked in order (more specific first).
_SECTIONS = [
    (re.compile(r"other\s+(income|revenue)", re.I), "Other Income"),
    (re.compile(r"other\s+expense", re.I), "Other Expense"),
    (re.compile(r"cost\s+of\s+(goods|sales|revenue|services)|direct\s+cost|\bcogs\b", re.I), "COGS"),
    (re.compile(r"^(ordinary\s+)?(income|revenue|revenues|sales|fee\s+income)\b", re.I), "Revenue"),
    (re.compile(r"expense|overhead|operating\s+cost", re.I), "Expense"),
]
_PREFIX_SECTIONS = {"Revenue": "Revenue", "COGS": "COGS", "Expense": "Expense"}


def _month_from_header(val, default_year: int | None) -> date | None:
    """First of the month a header cell names, or None."""
    if isinstance(val, datetime):
        return date(val.year, val.month, 1)
    if isinstance(val, date):
        return date(val.year, val.month, 1)
    if not isinstance(val, str):
        return None
    s = val.strip()
    m = _MONTH_NAME_RE.match(s)
    if m and m.group(1).lower() in _MONTHS:
        month = _MONTHS[m.group(1).lower()]
        yr = m.group(2)
        if yr:
            year = int(yr) + (2000 if len(yr) == 2 else 0)
        elif default_year:
            year = default_year
        else:
            return None
        return date(year, month, 1)
    m = _MONTH_NUM_RE.match(s)
    if m:
        if m.group(1):
            month, year = int(m.group(1)), int(m.group(2))
        else:
            year, month = int(m.group(3)), int(m.group(4))
        if 1 <= month <= 12:
            return date(year, month, 1)
    return None


def _month_end(d: date) -> date:
    return date(d.year, d.month, calendar.monthrange(d.year, d.month)[1])


def _number(val) -> float | None:
    if val is None or isinstance(val, bool):
        return None
    if isinstance(val, (int, float)):
        return float(val)
    if isinstance(val, str):
        s = val.strip()
        if not s or not re.search(r"\d", s):
            return None
        # Keep only things that look like money, not labels with digits in them.
        if not re.fullmatch(r"[\s$€£(),.\-\d]+(USD)?", s):
            return None
        return parse_money(s)
    return None


def _section_for(label: str) -> str | None:
    for rx, section in _SECTIONS:
        if rx.search(label):
            return section
    return None


def _find_header(ws, max_scan: int = 40):
    """(header_row, {col: month_date}, annual_col, title_year).

    Picks the row with the most month-like cells; with none, a row with a
    Total/YTD/year column instead."""
    title_years: list[int] = []
    for r in range(1, min(ws.max_row, max_scan) + 1):
        for c in range(1, min(ws.max_column, 6) + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str):
                title_years += [int(y) for y in _YEAR_RE.findall(v)]
    title_year = title_years[-1] if title_years else None

    best = (None, {}, None)
    for r in range(1, min(ws.max_row, max_scan) + 1):
        months = {}
        for c in range(1, ws.max_column + 1):
            d = _month_from_header(ws.cell(r, c).value, title_year)
            if d:
                months[c] = d
        if len(months) >= 2 and len(months) > len(best[1]):
            best = (r, months, None)
    if best[0]:
        return best[0], best[1], None, title_year

    for r in range(1, min(ws.max_row, max_scan) + 1):
        for c in range(2, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and _TOTAL_HEADER_RE.match(v.strip()):
                yrs = _YEAR_RE.findall(v)
                return r, {}, c, int(yrs[-1]) if yrs else title_year
            if isinstance(v, int) and 2000 <= v <= 2100:
                return r, {}, c, v
    return None, {}, None, title_year


def _entity_for(ws, header_row: int, path: Path, n_sheets: int) -> str:
    for r in range(1, header_row):
        for c in range(1, min(ws.max_column, 6) + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and _ENTITY_HINT_RE.search(v) and len(v) < 80:
                return v.strip()
    if n_sheets > 1:
        return ws.title.strip()
    return path.stem


def _parse_sheet(ws, path: Path, n_sheets: int) -> list[dict]:
    header_row, month_cols, annual_col, year = _find_header(ws)
    if header_row is None:
        return []
    amount_cols = sorted(month_cols) if month_cols else [annual_col]
    first_amount_col = min(amount_cols)
    if not month_cols and year is None:
        year = config.BUDGET_YEAR - 1
        logger.warning("%s / %s: no year found in the title; assuming %d", path.name, ws.title, year)

    entity = _entity_for(ws, header_row, path, n_sheets)
    is_consolidated = bool(_CONSOLIDATED_RE.search(ws.title)) and n_sheets > 1
    rows: list[dict] = []
    section: str | None = None
    groups: dict[int, str] = {}  # depth -> category header text

    for r in range(header_row + 1, ws.max_row + 1):
        label, depth = None, 0
        for c in range(1, first_amount_col):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip():
                label = v
                depth = (c - 1) * 4 + (len(v) - len(v.lstrip()))
                break
        if not label:
            continue
        text = label.strip()
        amounts = {c: _number(ws.cell(r, c).value) for c in amount_cols}
        has_amount = any(a is not None for a in amounts.values())

        if _SKIP_RE.match(text):
            continue
        if not has_amount:
            # A numbered header ("6000 · Payroll Expenses") is a parent
            # account, not a section, even if its name says "Expenses".
            sec = None if _ACCOUNT_RE.match(text) else _section_for(text)
            if sec and (section is None or depth <= min(groups, default=depth)):
                section = sec
                groups = {}
            else:
                groups = {d: g for d, g in groups.items() if d < depth}
                groups[depth] = text
            continue

        account_number, line_item = None, text
        m = _ACCOUNT_RE.match(text)
        if m:
            account_number, line_item = m.group(1), m.group(2).strip()
        row_section = section
        if row_section is None and account_number:
            row_section = _PREFIX_SECTIONS.get(config.GL_ACCOUNT_TYPE_PREFIXES.get(account_number[0], ""))
        if row_section is None:
            row_section = "Expense"
        parents = [g for d, g in sorted(groups.items()) if d < depth]
        category = parents[-1] if parents else None
        if category:
            cm = _ACCOUNT_RE.match(category)
            category = cm.group(2).strip() if cm else category

        for c, amt in amounts.items():
            if amt is None:
                continue
            if month_cols:
                start = month_cols[c]
                end, months = _month_end(start), 1
            else:
                start, end, months = date(year, 1, 1), date(year, 12, 31), 12
            rows.append({
                "entity": entity,
                "sheet": ws.title,
                "is_consolidated": is_consolidated,
                "section": row_section,
                "category": category,
                "account_number": account_number,
                "line_item": line_item,
                "period_start": start,
                "period_end": end,
                "months": months,
                "amount": amt,
                "source_file": path.name,
            })
    return rows


def parse(raw_dir: Path | None = None) -> list[dict]:
    raw_dir = raw_dir or config.RAW_DATA_DIR
    files = [f for f in find_all_files(raw_dir, config.SOURCE_FILE_PATTERNS["pl_workbook"])
             if not f.name.startswith("~$")]
    if not files:
        logger.info("No P&L workbook found in %s (patterns: %s)", raw_dir, config.SOURCE_FILE_PATTERNS["pl_workbook"])
        return []
    all_rows: list[dict] = []
    for path in files:
        try:
            wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
        except Exception as e:  # noqa: BLE001 -- one bad file shouldn't stop the build
            logger.warning("Skipping %s: %s", path.name, e)
            continue
        n_sheets = len(wb.worksheets)
        file_rows: list[dict] = []
        for ws in wb.worksheets:
            sheet_rows = _parse_sheet(ws, path, n_sheets)
            if sheet_rows:
                logger.info(
                    "P&L %s / %s: %d line items, %s to %s (entity %r)",
                    path.name, ws.title, len({r["line_item"] for r in sheet_rows}),
                    min(r["period_start"] for r in sheet_rows), max(r["period_end"] for r in sheet_rows),
                    sheet_rows[0]["entity"],
                )
            file_rows += sheet_rows
        if not file_rows:
            logger.warning("%s: no month or Total columns found -- check the layout", path.name)
        all_rows += file_rows

    # The same entity + period can appear in several files (e.g. re-exports).
    # Newest file wins, since files come back oldest first.
    latest: dict[tuple, str] = {}
    for r in all_rows:
        latest[(r["entity"], r["sheet"], r["period_start"], r["period_end"])] = r["source_file"]
    return [r for r in all_rows if latest[(r["entity"], r["sheet"], r["period_start"], r["period_end"])] == r["source_file"]]


def write_processed(rows: list[dict], out_path: Path | None = None) -> Path | None:
    if not rows:
        return None
    out_path = out_path or (config.PROCESSED_DATA_DIR / "pl_actuals.csv")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
    logger.info("Wrote %d rows -> %s", len(rows), out_path)
    return out_path


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s: %(message)s")
    out = parse()
    for section in ("Revenue", "COGS", "Expense", "Other Income", "Other Expense"):
        total = sum(r["amount"] for r in out if r["section"] == section)
        print(f"{section:<14} {total:>16,.2f}")
