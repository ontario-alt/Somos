"""
Parse the monthly billable tabs of the firm's "Billing Hours Analysis"
workbook (e.g. "DRAFT_Billing_Hours_Analysis.xlsx") -- tabs named like
"FY2025 — Monthly Billable" with one row per timekeeper, one column per
month ("Oct '24" ... "Sep '25"), grouped under "Attorneys" / "Planners" /
"Other Profs" header rows, then a TOTAL row and a small chart-data block.

This is the only source found so far for FY2025 (Oct 2024 - Sep 2025),
which the 2-year promotion lookback for FY2026 needs. It carries
*billable hours only* -- no credited / not-credited split -- so rows are
written to the same `monthly_hours` table as the firm's Monthly Hours
Report workbook with has_breakdown = False, and the page then scores
only the billable side (Total Activity can't be computed from it).

Where the Monthly Hours Report workbook already covers a month, that
source wins (build_warehouse drops these rows for those months): it has
the full breakdown, and this workbook's FY2024 tab is a copy of its
"Billable" column, which already includes credited non-billable hours.

The group header a row sits under is kept as `role_hint` (Attorneys ->
Attorney; Planners and Other Profs, i.e. Project Specialists -> Planner);
it applies only to people not listed in employee_targets.csv. Bracketed tags like "[Contractor]" or
"[Executive Office]" mark people without a role requirement.

Other tabs (the YTD "Billable Hours Analysis" summary, cost per billable
hour, profitability at A-rates) are derived figures, not source data,
and aren't loaded.
"""
from __future__ import annotations

import datetime
import logging
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import config
from etl.common import find_all_files, name_key

logger = logging.getLogger("somos.etl.billing_hours_analysis")

_MONTH_RE = re.compile(r"^(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s*['’]?\s*(\d{2,4})$", re.I)
_MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}
# "Other Profs" are Project Specialists (1,600, same as Planners). Anyone
# in these groups without a requirement (e.g. a contractor) is blank in
# employee_targets.csv, which overrides the group.
_GROUP_ROLES = {"attorneys": "Attorney", "planners": "Planner", "other profs": "Planner",
                "other professionals": "Planner"}
_TAG_RE = re.compile(r"\[([^\]]+)\]")


def _month(v) -> datetime.date | None:
    if isinstance(v, datetime.datetime):
        return v.date().replace(day=1)
    m = _MONTH_RE.match(str(v or "").strip())
    if not m:
        return None
    year = int(m.group(2))
    year = year + 2000 if year < 100 else year
    return datetime.date(year, _MONTHS[m.group(1)[:3].lower()], 1)


def _parse_sheet(ws, source_file: str) -> list[dict]:
    hdr_row, month_cols = None, {}
    for r in range(1, 6):
        cols = {c: _month(ws.cell(r, c).value) for c in range(2, ws.max_column + 1)}
        cols = {c: m for c, m in cols.items() if m}
        if len(cols) >= 6:
            hdr_row, month_cols = r, cols
            break
    if hdr_row is None:
        return []

    rows, group = [], None
    for r in range(hdr_row + 1, ws.max_row + 1):
        name = ws.cell(r, 1).value
        if not isinstance(name, str) or not name.strip():
            continue
        name = name.strip()
        values = {m: ws.cell(r, c).value for c, m in month_cols.items()}
        numeric = {m: float(v) for m, v in values.items() if isinstance(v, (int, float))}
        if name.upper().startswith("TOTAL"):
            break  # chart-data block below repeats the names
        if not numeric and all(v in (None, "") for v in values.values()):
            group = name.lower()  # "Attorneys" / "Planners" / "Other Profs"
            continue
        tag = _TAG_RE.search(name)
        clean = _TAG_RE.sub("", name).strip()
        role_hint = None if tag else _GROUP_ROLES.get(group)
        for month, hrs in numeric.items():
            rows.append(
                {
                    "employee_name": clean,
                    "name_key": name_key(clean),
                    "month": month,
                    "partial_month": False,
                    "billable_to_clients": hrs,
                    "nb_credited": 0.0,
                    "nb_not_credited": 0.0,
                    "total_hours": None,
                    "billable_requirement": None,
                    "credit_cap": None,
                    "total_requirement": None,
                    "note": f"[{tag.group(1)}]" if tag else None,
                    "sheet": ws.title,
                    "source_file": source_file,
                    "has_breakdown": False,
                    "role_hint": role_hint,
                }
            )
    return rows


def parse(paths: list[Path] | None = None) -> list[dict]:
    import openpyxl

    paths = paths if paths is not None else find_all_files(
        config.RAW_DATA_DIR, config.SOURCE_FILE_PATTERNS["billing_hours_analysis"]
    )
    by_key: dict[tuple, dict] = {}
    for path in paths:  # oldest -> newest; newer file wins per person-month
        logger.info("Parsing billing hours analysis workbook: %s", path)
        wb = openpyxl.load_workbook(path, data_only=True)
        for ws in wb.worksheets:
            if "monthly billable" not in ws.title.lower():
                continue
            rows = _parse_sheet(ws, path.name)
            logger.info("  %s: %d timekeeper-months, %d timekeepers", ws.title, len(rows),
                        len({r["name_key"] for r in rows}))
            for r in rows:
                by_key[(r["name_key"], r["month"])] = r
    return list(by_key.values())
