"""
Parse the Vantagepoint "Matter List" export (.xlsx) into a matter master
dimension table: matter code -> matter name -> client -> entity, plus an
"Organization Name" field (e.g. "LLC Planning", "LLP Legal", "LLC Admin")
that doubles as a practice-group/department taxonomy -- the piece that
was missing for quarterly's "revenue by matter type" and the Measuring
Period page's practice-group profitability.

Source shape: nested by (Company: <prefix> <entity> -> Billing Client
Name: <client> -> matter rows), flattened with no indentation, similar
in spirit to the AR/WIP exports but only 2 levels deep and with no
numeric subtotals to reconcile against -- there's nothing to sum here,
so correctness is checked by matter-code count instead (see __main__).

A client section with a blank "Billing Client Name:" (seen under every
entity) holds internal/admin matters -- office admin, PTO, holiday,
pro bono buckets -- not tied to a real client. These come through with
client_name = None rather than being dropped, since they're still real
matters that can show up in WIP/AR.

Output grain: one row per matter code:
    entity, client_name, matter_code, matter_name, organization_name,
    source_file
"""
from __future__ import annotations

import logging
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl

import config
from etl.common import find_latest_file

logger = logging.getLogger("somos.etl.matter_list")

_COMPANY_RE = re.compile(r"Company:\s*(\w+)\s+(.+)")
_CLIENT_RE = re.compile(r"Billing Client Name:\s*(.*)")


def parse(path: Path | None = None) -> list[dict]:
    path = path or find_latest_file(config.RAW_DATA_DIR, config.SOURCE_FILE_PATTERNS["matter_list"])
    if path is None:
        logger.warning(
            "No Matter List export found in %s -- skipping. Practice-group "
            "taxonomy and matter-master validation will be unavailable "
            "until one is added.",
            config.RAW_DATA_DIR,
        )
        return []
    logger.info("Parsing Matter List: %s", path)
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active

    rows: list[dict] = []
    entity = None
    client_name = None

    for r in range(1, ws.max_row + 1):
        col_a = ws.cell(r, 1).value
        col_a = col_a.strip() if isinstance(col_a, str) else col_a

        if col_a:
            m = _COMPANY_RE.match(col_a)
            if m:
                prefix, full_name = m.group(1), m.group(2).strip()
                entity = config.MATTER_CODE_ENTITY_PREFIXES.get(prefix, full_name)
                client_name = None
                continue
            m = _CLIENT_RE.match(col_a)
            if m:
                client_name = m.group(1).strip() or None
                continue
            if col_a.startswith("Total for") or col_a == "Matter List":
                continue

        matter_code = ws.cell(r, 3).value
        matter_name = ws.cell(r, 4).value
        organization_name = ws.cell(r, 5).value
        if not matter_code or matter_code == "Matter":  # the repeated header row
            continue
        rows.append(
            {
                "entity": entity,
                "client_name": client_name,
                "matter_code": matter_code.strip(),
                "matter_name": (matter_name or "").strip(),
                "organization_name": (organization_name or "").strip() or None,
                "source_file": path.name,
            }
        )

    logger.info("Parsed %d matters", len(rows))
    return rows


def write_processed(rows: list[dict], out_path: Path | None = None) -> Path | None:
    if not rows:
        logger.info("No matter list rows to write (source not available)")
        return None
    out_path = out_path or (config.PROCESSED_DATA_DIR / "matter_list.csv")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    import csv

    with open(out_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
    logger.info("Wrote %d rows -> %s", len(rows), out_path)
    return out_path


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s: %(message)s")
    rows = parse()
    write_processed(rows)
    print(f"{len(rows)} matters")
    by_entity = {}
    for r in rows:
        by_entity.setdefault(r["entity"], 0)
        by_entity[r["entity"]] += 1
    for e, n in by_entity.items():
        print(f"  {e}: {n} matters")
    n_no_client = sum(1 for r in rows if r["client_name"] is None)
    print(f"{n_no_client} matters with no client (internal/admin)")
    orgs = sorted({r["organization_name"] for r in rows if r["organization_name"]})
    print(f"{len(orgs)} distinct organization/practice-group values: {orgs[:10]}{'...' if len(orgs) > 10 else ''}")
