"""
Parse the firm's hand-built "Monthly Hours Report by Timekeeper" workbook
(e.g. "2024_Monthly_Hours_Report_by_Timekeeper.xlsx") -- the report the
measuring-period hours review was done in before this dashboard. Loading
it backfills closed measuring periods (FY24 = Oct 2023 - Sep 2024 in the
first sample) so they can be reviewed on the same page, and doubles as a
reconciliation check against the labor-detail-based numbers.

Source shape, checked against the sample: one tab per timekeeper.

    row 1: "Timekeeper" | "Billable Requirement" | "Creditable" | [blank] |
           "Total Combined Hours Required"     (label positions drift by a
           column between tabs -- located by label, not position)
    row 2: name | 1900 | 75 | [blank] | 2200   (requirements blank for
           people without a target, e.g. mid-year hires)
    row 3: column headers -- "Billable" (or "Total Billable") = billable to
           clients + credited non-billable, "Billable to Clients",
           "Non-Billable Hours Credited (Subject to Year Limit)",
           "Non-Billable Hours Not Credited", "Total Billable and
           Non-Billable"
    rows 4..15: one per month, month in column B as a date -- or as text
           for a partial month (e.g. "Sept 1-15" for someone who left
           mid-month), flagged partial_month
    then a totals row (blank column B), sometimes a free-text note.

Template tabs with no name in row 2 ("Sheet1", "New Sample") are skipped.
Note the report's own "Billable" column adds *all* credited non-billable
hours, uncapped -- the 75-hour year limit is applied on the dashboard,
not here, so this table keeps the raw monthly figures.

Output grain: one row per (timekeeper, month):
    employee_name, name_key, month, partial_month, billable_to_clients,
    nb_credited, nb_not_credited, total_hours, billable_requirement,
    credit_cap, total_requirement, note, sheet, source_file
"""
from __future__ import annotations

import csv
import datetime
import logging
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import config
from etl.common import find_all_files, name_key

logger = logging.getLogger("somos.etl.monthly_hours_workbook")

_MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], start=1)}


def _norm(v) -> str:
    return re.sub(r"[^a-z]", "", str(v or "").lower())


def _num(v) -> float | None:
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip()) if v not in (None, "") else None
    except ValueError:
        return None


def _month_of(v, prev: datetime.date | None) -> tuple[datetime.date | None, bool]:
    """(first-of-month date, partial) from a date cell or a label like
    'Sept 1-15'. Text labels take their year from the previous row."""
    if isinstance(v, datetime.datetime):
        return v.date().replace(day=1), False
    if isinstance(v, datetime.date):
        return v.replace(day=1), False
    s = str(v or "").strip().lower()
    if not s:
        return None, False
    m = _MONTHS.get(s[:3])
    if m is None:
        return None, False
    if prev is None:
        return None, False
    year = prev.year + (1 if m < prev.month else 0)
    return datetime.date(year, m, 1), True


def _header_positions(ws, max_row: int = 4) -> tuple[int | None, dict]:
    """Find the monthly column-header row and map each measure to a column."""
    for r in range(1, max_row + 1):
        labels = {c: _norm(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)}
        cols = {}
        for c, lab in labels.items():
            if lab.startswith("nonbillablehourscredited"):
                cols["nb_credited"] = c
            elif lab.startswith("nonbillablehoursnotcredited"):
                cols["nb_not_credited"] = c
            elif lab.startswith("billabletoclients"):
                cols["billable_to_clients"] = c
            elif lab.startswith("totalbillableandnonbillable"):
                cols["total_hours"] = c
            elif lab in ("billable", "totalbillable"):
                cols["billable_incl_credited"] = c
        if "nb_credited" in cols and ("billable_to_clients" in cols or "billable_incl_credited" in cols):
            return r, cols
    return None, {}


def _parse_sheet(ws, source_file: str) -> list[dict]:
    hdr_row, cols = _header_positions(ws)
    if hdr_row is None or hdr_row < 3:
        return []  # template tab without a name/requirements row

    # Requirement labels in row 1, values directly beneath in row 2.
    req = {}
    for c in range(1, ws.max_column + 1):
        lab = _norm(ws.cell(1, c).value)
        if lab.startswith("billablerequir"):
            req["billable_requirement"] = _num(ws.cell(2, c).value)
        elif lab == "creditable":
            req["credit_cap"] = _num(ws.cell(2, c).value)
        elif lab.startswith("totalcombined"):
            req["total_requirement"] = _num(ws.cell(2, c).value)
    name = str(ws.cell(2, 1).value or "").strip()
    if not name:
        return []

    note = None
    rows, prev = [], None
    for r in range(hdr_row + 1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and a.strip().lower().startswith("note"):
            note = a.strip()
            continue
        month, partial = _month_of(ws.cell(r, 2).value, prev)
        if month is None:
            continue  # totals row / blank
        prev = month
        g = {k: _num(ws.cell(r, c).value) for k, c in cols.items()}
        credited = g.get("nb_credited") or 0.0
        billable = g.get("billable_to_clients")
        if billable is None and g.get("billable_incl_credited") is not None:
            billable = g["billable_incl_credited"] - credited
        billable = billable or 0.0
        not_credited = g.get("nb_not_credited") or 0.0
        total = g.get("total_hours")
        rows.append(
            {
                "employee_name": name,
                "name_key": name_key(name),
                "month": month,
                "partial_month": partial,
                "billable_to_clients": billable,
                "nb_credited": credited,
                "nb_not_credited": not_credited,
                "total_hours": total if total is not None else billable + credited + not_credited,
                "billable_requirement": req.get("billable_requirement"),
                "credit_cap": req.get("credit_cap"),
                "total_requirement": req.get("total_requirement"),
                "note": None,
                "sheet": ws.title,
                "source_file": source_file,
                "has_breakdown": True,
                "role_hint": None,
            }
        )
    if note:
        for row in rows:
            row["note"] = note
    return rows


def _parse_one(path: Path) -> list[dict]:
    import openpyxl

    logger.info("Parsing monthly hours workbook: %s", path)
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    out = []
    for ws in wb.worksheets:
        rows = _parse_sheet(ws, path.name)
        if rows:
            out.extend(rows)
        else:
            logger.info("  skipped tab %r (template / no timekeeper name)", ws.title)
    names = sorted({r["employee_name"] for r in out})
    logger.info("Parsed %d timekeeper-months for %d timekeepers from %s", len(out), len(names), path.name)
    return out


def parse(paths: list[Path] | None = None) -> list[dict]:
    """Every matching workbook is read (one per measuring period). If two
    cover the same person-month, the newer file wins."""
    paths = paths if paths is not None else find_all_files(
        config.RAW_DATA_DIR, config.SOURCE_FILE_PATTERNS["monthly_hours_workbook"]
    )
    by_key: dict[tuple, dict] = {}
    for path in paths:  # oldest -> newest
        for r in _parse_one(path):
            by_key[(r["name_key"], r["month"])] = r
    return list(by_key.values())


def write_processed(rows: list[dict], out_path: Path | None = None) -> Path | None:
    if not rows:
        logger.info("No monthly hours workbook rows to write (source not available)")
        return None
    out_path = out_path or (config.PROCESSED_DATA_DIR / "monthly_hours.csv")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
    logger.info("Wrote %d rows -> %s", len(rows), out_path)
    return out_path


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s: %(message)s")
    rows = parse()
    write_processed(rows)
    print(f"{len(rows)} timekeeper-months")
