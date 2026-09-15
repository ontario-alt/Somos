"""
Parse the hand-built "Somos_AR_Aging_MM.DD.YYYY.xlsx" workbook -- the
multi-tab report (Summary / Week Over Week / AR Aging Detail / Priority
Board / Client Rollup) that defined this app's own weekly AR page format
in the first place. Unlike the raw Vantagepoint "All AR Report" export
(parse_ar_detail.py), this workbook's "AR Aging Detail" tab already has
real Entity/Client/Matter columns with real headers -- no "Total for
<Client> - <Matter>" string-splitting or PDF position-matching needed.

Feeds the same ar_aging_detail table as parse_ar_detail.py (same grain,
same columns), so historical copies of this workbook are a second, easier
way to backfill AR aging history -- each file's own "Aged as of" date
(read from the tab's own subtitle line) becomes its own snapshot.

Output grain: one row per (entity, client, matter):
    entity, client_name, matter_name, client_name_confidence,
    current_0_30, days_31_60, days_61_90, days_91_120, over_120,
    balance, as_of_date, source_file
"""
from __future__ import annotations

import logging
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import config
from etl.common import find_all_files, parse_vp_date

logger = logging.getLogger("somos.etl.ar_aging_workbook")

_SHEET_NAME = "AR Aging Detail"
_AS_OF_RE = re.compile(r"Aged as of\s+(\d{1,2}/\d{1,2}/\d{4})")
_HEADER_MAP = {
    "Entity": "entity",
    "Client": "client_name",
    "Matter": "matter_name",
    "Current": "current_0_30",
    "31-60": "days_31_60",
    "61-90": "days_61_90",
    "91-120": "days_91_120",
    "Over 120": "over_120",
    "Balance": "balance",
}


def _parse_one(path: Path) -> list[dict]:
    logger.info("Parsing AR Aging workbook: %s", path)
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    if _SHEET_NAME not in wb.sheetnames:
        logger.warning("%s has no '%s' tab -- skipping", path.name, _SHEET_NAME)
        return []
    ws = wb[_SHEET_NAME]

    as_of = None
    header_row = None
    col_idx: dict[str, int] = {}
    for r in range(1, min(ws.max_row, 10) + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if not isinstance(v, str):
                continue
            if as_of is None:
                m = _AS_OF_RE.search(v)
                if m:
                    as_of = parse_vp_date(m.group(1))
            if v.strip() in _HEADER_MAP:
                col_idx[_HEADER_MAP[v.strip()]] = c
        if {"entity", "client_name", "matter_name", "balance"} <= col_idx.keys():
            header_row = r
            break

    if header_row is None:
        logger.warning("Couldn't find the header row in %s's '%s' tab -- skipping", path.name, _SHEET_NAME)
        return []

    rows: list[dict] = []
    for r in range(header_row + 1, ws.max_row + 1):
        entity_raw = ws.cell(r, col_idx["entity"]).value
        client_name = ws.cell(r, col_idx["client_name"]).value
        matter_name = ws.cell(r, col_idx["matter_name"]).value
        if not client_name or not matter_name:
            continue
        rows.append(
            {
                "entity": config.MATTER_CODE_ENTITY_PREFIXES.get(str(entity_raw).strip(), entity_raw),
                "client_name": str(client_name).strip(),
                "matter_name": str(matter_name).strip(),
                "client_name_confidence": "ok",
                "as_of_date": as_of,
                "source_file": path.name,
                "current_0_30": round(ws.cell(r, col_idx["current_0_30"]).value or 0.0, 2),
                "days_31_60": round(ws.cell(r, col_idx["days_31_60"]).value or 0.0, 2),
                "days_61_90": round(ws.cell(r, col_idx["days_61_90"]).value or 0.0, 2),
                "days_91_120": round(ws.cell(r, col_idx["days_91_120"]).value or 0.0, 2),
                "over_120": round(ws.cell(r, col_idx["over_120"]).value or 0.0, 2),
                "balance": round(ws.cell(r, col_idx["balance"]).value or 0.0, 2),
            }
        )

    total = round(sum(r["balance"] for r in rows), 2)
    logger.info("Parsed %d client/matter AR rows from %s (as of %s, total %.2f)", len(rows), path.name, as_of, total)
    return rows


def parse(paths: list[Path] | None = None) -> list[dict]:
    """Parses every matching workbook in data/raw/, not just the newest --
    each carries its own "Aged as of" date, so dropping several weeks in
    at once backfills real history in a single build_warehouse.py run."""
    paths = paths if paths is not None else find_all_files(config.RAW_DATA_DIR, config.SOURCE_FILE_PATTERNS["ar_aging_workbook"])
    if not paths:
        return []
    all_rows: list[dict] = []
    for path in paths:
        all_rows.extend(_parse_one(path))
    return all_rows


def write_processed(rows: list[dict], out_path: Path | None = None) -> Path | None:
    if not rows:
        logger.info("No AR Aging workbook rows to write (source not available)")
        return None
    out_path = out_path or (config.PROCESSED_DATA_DIR / "ar_aging_workbook.csv")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    import csv

    with open(out_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
    logger.info("Wrote %d rows -> %s", len(rows), out_path)
    return out_path


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s: %(message)s")
    rows = parse()
    write_processed(rows)
    print(f"{len(rows)} client/matter AR rows across {len({r['as_of_date'] for r in rows})} snapshot date(s)")
