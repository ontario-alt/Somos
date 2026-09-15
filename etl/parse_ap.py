"""
Parse the Vantagepoint AP export (.xlsx) into two tidy tables: open AP
aging (unpaid vendor invoices, bucketed like AR) and cash disbursements
(the payment lines in the same export -- "AP Disb"/"Auto Check" rows
with a Check Date, i.e. money that actually went out the door).

Source shape: already flat (no nested tree, unlike AR/WIP). Each vendor
invoice shows up as one or more "bill" rows (one per matter/phase split,
each with a positive Voucher Amount and Balance) followed by matching
payment rows (Description "AP Disb" or "Auto Check", a Check Date, and a
negative Balance that offsets the bill). Grouping by (Vendor Number,
Invoice Number) and summing Balance nets to ~0 for a fully paid invoice
and to the outstanding amount for an open one -- there are no separate
"paid" vs "open" flags in the export, so this net-balance approach is
the only way to tell them apart.

Entity is derived per invoice group from whichever of these resolves
first: the matter code prefix (LLC/LLP, same convention as AR/WIP) on
its bill line, then the Liability Code ("PWP-LLC"/"PWP-LLP"), then the
Voucher BankCode ("LLC"/"LLP"/"BILLLLC"/"BILLLLP") -- payment lines and
overhead vouchers (placeholder matter "ZZZ00-100") have no usable matter
code, hence the fallback chain.

Output grain:
  ap_aging: one row per open (unpaid) vendor invoice:
    vendor_name, invoice_number, invoice_date, entity, balance,
    current_0_30, days_31_60, days_61_90, days_91_120, over_120,
    source_file
  cash_disbursements: one row per payment line:
    vendor_name, check_date, check_ref_no, amount, entity, source_file
"""
from __future__ import annotations

import datetime
import logging
import re
import sys
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl

import config
from etl.common import find_latest_file

logger = logging.getLogger("somos.etl.ap")

_MATTER_PREFIX_RE = re.compile(r"^[A-Z]+")

_COLS = [
    "_blank",
    "vendor_number",
    "vendor_name",
    "invoice_date",
    "invoice_number",
    "voucher_date",
    "voucher_number",
    "pay_terms",
    "payment_date",
    "liability_code",
    "voucher_bank_code",
    "description",
    "matter",
    "phase",
    "account_number",
    "check_ref_no",
    "check_date",
    "voucher_amount",
    "previous_payments",
    "discount_taken",
    "balance",
]


def _entity_for_row(matter: str | None, liability_code: str | None, bank_code: str | None) -> str | None:
    if matter:
        m = _MATTER_PREFIX_RE.match(matter.strip())
        if m and m.group(0) in config.MATTER_CODE_ENTITY_PREFIXES:
            return config.MATTER_CODE_ENTITY_PREFIXES[m.group(0)]
    for candidate in (liability_code, bank_code):
        if not candidate:
            continue
        for suffix, entity in config.AP_ENTITY_CODE_SUFFIXES.items():
            if suffix in candidate:
                return entity
    return None


def _load_rows(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = []
    for r in range(2, ws.max_row + 1):
        raw = [ws.cell(r, c).value for c in range(1, len(_COLS) + 1)]
        row = dict(zip(_COLS, raw))
        if not row.get("vendor_number"):
            continue
        rows.append(row)
    return rows


def _bucket_for_age(days: int) -> str:
    if days <= 30:
        return config.AGING_BUCKETS[0]
    if days <= 60:
        return config.AGING_BUCKETS[1]
    if days <= 90:
        return config.AGING_BUCKETS[2]
    if days <= 120:
        return config.AGING_BUCKETS[3]
    return config.AGING_BUCKETS[4]


def parse(path: Path | None = None, as_of: datetime.date | None = None) -> tuple[list[dict], list[dict]]:
    path = path or find_latest_file(config.RAW_DATA_DIR, config.SOURCE_FILE_PATTERNS["ap_aging"])
    if path is None:
        logger.warning(
            "No AP export found in %s matching %r -- skipping. AP aging and "
            "cash-disbursement pages will be unavailable until one is added.",
            config.RAW_DATA_DIR,
            config.SOURCE_FILE_PATTERNS["ap_aging"],
        )
        return [], []
    as_of = as_of or datetime.date.today()
    logger.info("Parsing AP export: %s", path)
    rows = _load_rows(path)

    # Pass 1: group by invoice to net open balances and resolve each
    # invoice's entity from whichever of its rows has a usable signal --
    # a payment line usually has none of its own (blank matter, no
    # liability/bank code), but its sibling bill line does.
    groups: dict[tuple, dict] = {}
    for row in rows:
        key = (row["vendor_number"], row["invoice_number"])
        g = groups.setdefault(
            key,
            {
                "vendor_name": row["vendor_name"],
                "invoice_date": row["invoice_date"],
                "balance": 0.0,
                "entity": None,
            },
        )
        g["balance"] += float(row["balance"] or 0)
        entity = _entity_for_row(row.get("matter"), row.get("liability_code"), row.get("voucher_bank_code"))
        if entity and not g["entity"]:
            g["entity"] = entity

    # Pass 2: emit one row per payment line, using its invoice group's
    # resolved entity rather than trying to re-resolve from the row alone.
    disbursements = []
    for row in rows:
        check_date = row.get("check_date")
        if not (check_date and row.get("previous_payments")):
            continue
        key = (row["vendor_number"], row["invoice_number"])
        disbursements.append(
            {
                "vendor_name": row["vendor_name"],
                "check_date": check_date.date() if isinstance(check_date, datetime.datetime) else check_date,
                "check_ref_no": (row.get("check_ref_no") or "").strip() or None,
                "amount": float(row["previous_payments"]),
                "entity": groups[key]["entity"],
                "source_file": path.name,
            }
        )

    ap_aging = []
    for (vendor_number, invoice_number), g in groups.items():
        if abs(g["balance"]) < 0.01:
            continue  # fully paid, net zero -- not an open item
        invoice_date = g["invoice_date"]
        invoice_date = invoice_date.date() if isinstance(invoice_date, datetime.datetime) else invoice_date
        age_days = (as_of - invoice_date).days if invoice_date else None
        bucket = _bucket_for_age(age_days) if age_days is not None else None
        rec = {
            "vendor_name": g["vendor_name"],
            "invoice_number": invoice_number,
            "invoice_date": invoice_date,
            "entity": g["entity"],
            "balance": round(g["balance"], 2),
            "source_file": path.name,
        }
        for b in config.AGING_BUCKETS:
            rec[b] = round(g["balance"], 2) if b == bucket else 0.0
        ap_aging.append(rec)

    logger.info(
        "Parsed %d open AP invoices (%d total groups) and %d disbursement lines",
        len(ap_aging),
        len(groups),
        len(disbursements),
    )
    return ap_aging, disbursements


def write_processed(ap_aging: list[dict], disbursements: list[dict]) -> tuple[Path | None, Path | None]:
    config.PROCESSED_DATA_DIR.mkdir(parents=True, exist_ok=True)
    import csv

    ap_path = None
    if ap_aging:
        ap_path = config.PROCESSED_DATA_DIR / "ap_aging.csv"
        with open(ap_path, "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=list(ap_aging[0].keys()))
            w.writeheader()
            w.writerows(ap_aging)
        logger.info("Wrote %d rows -> %s", len(ap_aging), ap_path)

    disb_path = None
    if disbursements:
        disb_path = config.PROCESSED_DATA_DIR / "cash_disbursements.csv"
        with open(disb_path, "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=list(disbursements[0].keys()))
            w.writeheader()
            w.writerows(disbursements)
        logger.info("Wrote %d rows -> %s", len(disbursements), disb_path)

    return ap_path, disb_path


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s: %(message)s")
    ap_aging, disbursements = parse()
    write_processed(ap_aging, disbursements)
    print(f"{len(ap_aging)} open AP invoices, total {sum(r['balance'] for r in ap_aging):,.2f}")
    print(f"{len(disbursements)} disbursement lines, total {sum(r['amount'] for r in disbursements):,.2f}")
