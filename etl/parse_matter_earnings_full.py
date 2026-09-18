"""
Parse the Vantagepoint "Matter Earnings" export (.xlsx) -- a different,
much broader report than the "NTE Tracking Report" parse_matter_earnings.py
already handles. That one only lists matters with a not-to-exceed cap
set (~42 of ~315 matters firm-wide); this one lists every matter with
any JTD activity, client and overhead alike (163 matters in the sample
provided: 84 LLP + 30 LLC client matters, 43 LLC/LLP overhead buckets,
6 admin/PTO/sick/MX rows) -- a real, much less partial source for
billed-by-matter than matter_earnings has had until now.

Like the NTE Tracking Report and the Employee Cost Rate Details export,
this can be run "as" either active company but is actually one
firm-wide report (checked directly: matter codes span LLC/LLP/overhead/
MX prefixes regardless of which company header the export shows).

Source shape: flat, no phase/task sub-rows (unlike the companion Matter
List export sharing this matter universe) -- one row per matter, with
matter code and matter name concatenated in the first column (e.g.
"LLC25-002 PG&E Potter Valley Project Surrender") rather than split into
separate columns, and no "Final Totals" row to reconcile against.

Output grain: one row per matter code, in the SAME shape
parse_matter_earnings.py produces (so both sources merge into one
matter_earnings table in build_warehouse.py) --
    entity, matter_code, matter_name, client_name, charge_type,
    project_manager, nte, amount_left, wip_unbilled,
    amount_left_with_wip, jtd_hours, invoiced_to_date, jtd_revenue,
    jtd_profit, source_file
client_name, nte, amount_left, wip_unbilled, and amount_left_with_wip
have no equivalent column in this report and are always None -- this
report's value is JTD Billed (-> invoiced_to_date), JTD Revenue, and
JTD Profit for a far larger slice of the portfolio, not those fields.
"""
from __future__ import annotations

import logging
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import config
from etl.common import find_all_files

logger = logging.getLogger("somos.etl.matter_earnings_full")

_CODE_NAME_RE = re.compile(r"^(\S+)\s+(.*)$")
_PREFIX_RE = re.compile(r"^[A-Z]+")

_HEADER_ROW_MARKER = "Charge"  # matches the "Charge\nType" header cell


def _entity_for_matter_code(matter_code: str) -> str | None:
    m = _PREFIX_RE.match(matter_code)
    return config.MATTER_CODE_ENTITY_PREFIXES.get(m.group(0)) if m else None


def _parse_one(path: Path) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active

    header_row = None
    for r in range(1, min(ws.max_row, 20) + 1):
        if any(isinstance(ws.cell(r, c).value, str) and _HEADER_ROW_MARKER in ws.cell(r, c).value for c in range(1, ws.max_column + 1)):
            header_row = r
            break
    if header_row is None:
        logger.warning("Could not find the header row in %s -- skipping", path)
        return []

    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        raw = ws.cell(r, 1).value
        if not isinstance(raw, str) or not raw.strip():
            continue
        m = _CODE_NAME_RE.match(raw.strip())
        if not m:
            continue
        matter_code, matter_name = m.group(1), m.group(2).strip()
        rows.append(
            {
                "entity": _entity_for_matter_code(matter_code),
                "matter_code": matter_code,
                "matter_name": matter_name,
                "client_name": None,
                "charge_type": (ws.cell(r, 3).value or "").strip() or None,
                "project_manager": (ws.cell(r, 5).value or "").strip() or None,
                "nte": None,
                "amount_left": None,
                "wip_unbilled": None,
                "amount_left_with_wip": None,
                "jtd_hours": float(ws.cell(r, 8).value or 0),
                "invoiced_to_date": float(ws.cell(r, 10).value or 0),
                "jtd_revenue": float(ws.cell(r, 12).value or 0),
                "jtd_profit": float(ws.cell(r, 13).value or 0),
                "source_file": path.name,
            }
        )
    logger.info("Parsed %d matters from %s", len(rows), path.name)
    return rows


def parse(paths: list[Path] | None = None) -> list[dict]:
    paths = paths or find_all_files(config.RAW_DATA_DIR, config.SOURCE_FILE_PATTERNS["matter_earnings_full"])
    if not paths:
        logger.warning(
            "No Matter Earnings export found in %s -- skipping. Billed-by-matter coverage stays limited to "
            "the NTE Tracking Report's ~42-matter subset.",
            config.RAW_DATA_DIR,
        )
        return []

    by_code: dict[str, dict] = {}
    for path in paths:
        for row in _parse_one(path):
            existing = by_code.get(row["matter_code"])
            if existing and existing["invoiced_to_date"] != row["invoiced_to_date"]:
                logger.warning(
                    "Matter %s disagrees between %s and %s (JTD Billed %.2f vs %.2f) -- keeping the later file",
                    row["matter_code"],
                    existing["source_file"],
                    row["source_file"],
                    existing["invoiced_to_date"],
                    row["invoiced_to_date"],
                )
            by_code[row["matter_code"]] = row

    rows = list(by_code.values())
    logger.info("Parsed %d distinct matters across %d file(s)", len(rows), len(paths))
    return rows


def write_processed(rows: list[dict], out_path: Path | None = None) -> Path | None:
    if not rows:
        logger.info("No matter earnings (full) rows to write (source not available)")
        return None
    out_path = out_path or (config.PROCESSED_DATA_DIR / "matter_earnings_full.csv")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    import csv

    with open(out_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
    logger.info("Wrote %d rows -> %s", len(rows), out_path)
    return out_path


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s: %(message)s")
    rows = parse()
    write_processed(rows)
    print(f"{len(rows)} matters")
    by_entity: dict[str, int] = {}
    for r in rows:
        by_entity.setdefault(r["entity"], 0)
        by_entity[r["entity"]] += 1
    for e, n in by_entity.items():
        print(f"  {e}: {n} matters")
