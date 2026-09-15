"""
Parse the Vantagepoint "NTE Tracking Report" (Based On: Matter Earnings)
into a tidy matter-level earnings table -- real JTD revenue and profit
per matter, keyed by the same matter code AR/WIP/the matter list use.
This is what makes it possible to dollarize the origination credit
matrix (previously blocked: free-text matter names only matched AR's
~16% of the time) and to compute practice-group profitability for real
instead of using AR balance as a revenue proxy.

Like the Employee Cost Rate Details export, this report can be run
"as" either active company, but checked directly against both sample
files: identical content in both, matter-for-matter -- one firm-wide
report, not per-entity data. Deduplicated across files by matter code,
same approach as parse_employee_cost.py, with a warning if a future
export ever genuinely disagrees between files.

Source shape: two-level, no true nesting -- a "Matter Number: <code>
<name>" row followed by that matter's phase/task rows (a numeric phase
code like "0001 <phase name>"), each with the same figures scoped to
that phase. This parser keeps only the matter-level rows (the phase
breakdown is a finer grain than the app currently needs); a
"Final Totals" row at the end is used to verify the matter-level rows
reconcile.

Output grain: one row per matter code:
    entity, matter_code, matter_name, client_name, charge_type,
    project_manager, nte, amount_left, wip_unbilled,
    amount_left_with_wip, jtd_hours, invoiced_to_date, jtd_revenue,
    jtd_profit, source_file
"""
from __future__ import annotations

import logging
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import config
from etl.common import find_all_files

logger = logging.getLogger("somos.etl.matter_earnings")

_MATTER_RE = re.compile(r"^Matter Number:\s*(\S+)\s+(.*)$")
_PREFIX_RE = re.compile(r"^[A-Z]+")


def _entity_for_matter_code(matter_code: str) -> str | None:
    m = _PREFIX_RE.match(matter_code)
    return config.MATTER_CODE_ENTITY_PREFIXES.get(m.group(0)) if m else None


def _parse_one(path: Path) -> tuple[list[dict], dict | None]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active

    rows = []
    final_totals = None
    for r in range(1, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if not isinstance(label, str) or not label.strip():
            continue
        label = label.strip()

        if label == "Final Totals":
            final_totals = {
                "jtd_revenue": ws.cell(r, 16).value or 0.0,
                "jtd_profit": ws.cell(r, 17).value or 0.0,
            }
            continue

        m = _MATTER_RE.match(label)
        if not m:
            continue  # phase/task sub-row, not a matter-level row
        matter_code, matter_name = m.group(1), m.group(2).strip()
        rows.append(
            {
                "entity": _entity_for_matter_code(matter_code),
                "matter_code": matter_code,
                "matter_name": matter_name,
                "client_name": (ws.cell(r, 3).value or "").strip() or None,
                "charge_type": (ws.cell(r, 6).value or "").strip() or None,
                "project_manager": (ws.cell(r, 7).value or "").strip() or None,
                "nte": float(ws.cell(r, 8).value or 0),
                "amount_left": float(ws.cell(r, 9).value or 0),
                "wip_unbilled": float(ws.cell(r, 10).value or 0),
                "amount_left_with_wip": float(ws.cell(r, 11).value or 0),
                "jtd_hours": float(ws.cell(r, 13).value or 0),
                "invoiced_to_date": float(ws.cell(r, 14).value or 0),
                "jtd_revenue": float(ws.cell(r, 16).value or 0),
                "jtd_profit": float(ws.cell(r, 17).value or 0),
                "source_file": path.name,
            }
        )

    if final_totals:
        computed_revenue = round(sum(r["jtd_revenue"] for r in rows), 2)
        computed_profit = round(sum(r["jtd_profit"] for r in rows), 2)
        if abs(computed_revenue - final_totals["jtd_revenue"]) > 0.02 or abs(computed_profit - final_totals["jtd_profit"]) > 0.02:
            logger.warning(
                "Matter-level rows don't reconcile to Final Totals in %s: revenue %.2f vs %.2f, profit %.2f vs %.2f",
                path.name,
                computed_revenue,
                final_totals["jtd_revenue"],
                computed_profit,
                final_totals["jtd_profit"],
            )
        else:
            logger.info("Matter-level rows reconcile exactly to Final Totals in %s.", path.name)

    logger.info("Parsed %d matters from %s", len(rows), path.name)
    return rows, final_totals


def parse(paths: list[Path] | None = None) -> list[dict]:
    paths = paths or find_all_files(config.RAW_DATA_DIR, config.SOURCE_FILE_PATTERNS["nte_tracking"])
    if not paths:
        logger.warning(
            "No NTE Tracking Report found in %s -- skipping. Matter-level "
            "revenue/profit (and dollarized originations, real practice-group "
            "profitability) will be unavailable until one is added.",
            config.RAW_DATA_DIR,
        )
        return []

    by_matter: dict[str, dict] = {}
    conflicts: list[str] = []
    n_overlap = 0
    for path in paths:
        rows, _ = _parse_one(path)
        for row in rows:
            key = row["matter_code"]
            existing = by_matter.get(key)
            if existing:
                n_overlap += 1
                if existing["jtd_revenue"] != row["jtd_revenue"] or existing["jtd_profit"] != row["jtd_profit"]:
                    conflicts.append(f"{row['matter_code']}: {existing['source_file']} vs {row['source_file']}")
            by_matter[key] = row
    if conflicts:
        logger.warning(
            "%d matter(s) have different figures across NTE export files -- these may "
            "genuinely be entity-specific after all; verify before trusting them: %s",
            len(conflicts),
            conflicts,
        )
    elif n_overlap:
        logger.info(
            "%d matter(s) appeared in more than one export file, all with matching figures "
            "-- treated as one firm-wide report, not per-entity duplicates.",
            n_overlap,
        )
    return list(by_matter.values())


def write_processed(rows: list[dict], out_path: Path | None = None) -> Path | None:
    if not rows:
        logger.info("No matter earnings rows to write (source not available)")
        return None
    out_path = out_path or (config.PROCESSED_DATA_DIR / "matter_earnings.csv")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    import csv

    with open(out_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
    logger.info("Wrote %d rows -> %s", len(rows), out_path)
    return out_path


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s: %(message)s")
    rows = parse()
    write_processed(rows)
    print(f"{len(rows)} matters, total JTD revenue {sum(r['jtd_revenue'] for r in rows):,.2f}, total JTD profit {sum(r['jtd_profit'] for r in rows):,.2f}")
