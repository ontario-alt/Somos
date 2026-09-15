"""
Parse the Vantagepoint GL trial balance export (.xlsx) into a tidy table.

Source shape: unlike AR/WIP, this is already flat -- one row per GL
account, no nested tree. Layout (verified against Trial_Balance.xlsx):
    B6: entity name (e.g. "Somos Group LLC")
    C5: "For the period 9/1/2026 - 9/30/2026"
    Row 9: column headers (Opening Balance, Debits, Credits, Closing Balance)
    Rows 10+: "<account number>          <account name>" in column A
              (fixed-width padded, not delimited), then D/E/F/I = opening/
              debits/credits/closing balance
    "Subtotal" rows interspersed after each account block, and a
    trailing "Final Totals" row -- both skipped; account type (Asset/
    Liability/Equity/Revenue/COGS/Expense) is derived from the leading
    digit of the account number instead of trusting the report's own
    ad hoc subtotal groupings.

Vantagepoint trial balances are run per entity, so a real monthly drop
is likely to include one file per entity all matching the same glob
pattern -- parse() reads every matching file, not just the newest.

Output grain: one row per (entity, account) for the period covered:
    entity, period_start, period_end, account_number, account_name,
    account_type, opening_balance, debits, credits, closing_balance,
    source_file
"""
from __future__ import annotations

import logging
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl

import config
from etl.common import find_all_files, parse_vp_date

logger = logging.getLogger("somos.etl.gl")

_ACCOUNT_RE = re.compile(r"^(\d+)\s+(.*)$")
_PERIOD_RE = re.compile(r"For the period\s+(\d{1,2}/\d{1,2}/\d{4})\s*-\s*(\d{1,2}/\d{1,2}/\d{4})")


def _account_type(account_number: str) -> str | None:
    if not account_number:
        return None
    return config.GL_ACCOUNT_TYPE_PREFIXES.get(account_number[0])


def _parse_one(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active

    entity = None
    period_start = period_end = None
    header_row = None
    for r in range(1, min(ws.max_row, 20) + 1):
        for c in range(1, ws.max_column + 1):
            val = ws.cell(r, c).value
            if not isinstance(val, str):
                continue
            m = _PERIOD_RE.search(val)
            if m:
                period_start = parse_vp_date(m.group(1))
                period_end = parse_vp_date(m.group(2))
            if val.strip() == "" and c == 4 and header_row is None:
                pass
        header_val = ws.cell(r, 4).value
        if isinstance(header_val, str) and "Opening" in header_val:
            header_row = r
        # Entity name: first non-empty text cell in column B before the header row.
        b_val = ws.cell(r, 2).value
        if entity is None and isinstance(b_val, str) and b_val.strip() and "For the period" not in b_val:
            entity = b_val.strip()

    if header_row is None:
        raise ValueError(f"Couldn't find the Opening/Debits/Credits header row in {path}")
    if entity is None:
        raise ValueError(f"Couldn't find the entity name (expected in column B) in {path}")

    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if not isinstance(label, str) or not label.strip():
            continue
        label = label.strip()
        if label in ("Subtotal", "Final Totals"):
            continue
        m = _ACCOUNT_RE.match(label)
        if not m:
            logger.warning("Skipping unrecognized GL row in %s: %r", path, label)
            continue
        account_number, account_name = m.group(1), m.group(2).strip()
        opening = ws.cell(r, 4).value
        debits = ws.cell(r, 5).value
        credits_ = ws.cell(r, 6).value
        closing = ws.cell(r, 9).value
        rows.append(
            {
                "entity": entity,
                "period_start": period_start,
                "period_end": period_end,
                "account_number": account_number,
                "account_name": account_name,
                "account_type": _account_type(account_number),
                "opening_balance": float(opening or 0),
                "debits": float(debits or 0),
                "credits": float(credits_ or 0),
                "closing_balance": float(closing or 0),
                "source_file": path.name,
            }
        )
    logger.info("Parsed %d GL accounts for %s from %s", len(rows), entity, path.name)
    return rows


def parse(paths: list[Path] | None = None) -> list[dict]:
    paths = paths or find_all_files(config.RAW_DATA_DIR, config.SOURCE_FILE_PATTERNS["gl_trial_balance"])
    if not paths:
        logger.warning(
            "No GL trial balance export found in %s matching %r -- skipping. "
            "P&L and profitability pages will be unavailable until one is added.",
            config.RAW_DATA_DIR,
            config.SOURCE_FILE_PATTERNS["gl_trial_balance"],
        )
        return []
    rows: list[dict] = []
    for path in paths:
        rows.extend(_parse_one(path))
    return rows


def write_processed(rows: list[dict], out_path: Path | None = None) -> Path | None:
    if not rows:
        logger.info("No GL rows to write (source not available)")
        return None
    out_path = out_path or (config.PROCESSED_DATA_DIR / "gl_trial_balance.csv")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    import csv

    with open(out_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
    logger.info("Wrote %d rows -> %s", len(rows), out_path)
    return out_path


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s: %(message)s")
    rows = parse()
    write_processed(rows)
    print(f"{len(rows)} GL account rows")
    if rows:
        by_type = {}
        for r in rows:
            by_type.setdefault(r["account_type"], 0.0)
            by_type[r["account_type"]] += r["closing_balance"]
        for t, v in by_type.items():
            print(f"  {t}: {v:,.2f}")
